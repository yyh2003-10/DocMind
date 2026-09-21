"""P0 止血回归：坏 PPT 三症状（页码标题 / Quote 空壳 / 门禁）。

对应坏文件症状：
1. 页标题为「第 N 页」
2. Quote 页因运算符优先级 bug 丢真实金句
3. 空正文页 / 残留 markdown 仍被导出
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from doc2mind.core.creator import export_artifact, extract_artifact, validate_pptx_export
from doc2mind.core.creator.exporters.pptx_exporter import PptxExporter
from doc2mind.core.creator.models import ArtifactModel, ArtifactType, SlideLayoutType, SlideModel


# 复现坏 pptx 的等价 markdown：无逐页 # 标题，仅 artifact 标题 + 内容
BAD_MARP = """
:::artifact type="pptx" title="结构刚度关键指标——挠度（Deflection）深度解析" theme="tech_blue"
---
## 结构刚度关键指标——挠度（Deflection）深度解析
基于 DocMind 智能知识库生成
---
挠度是梁在荷载下竖向位移的度量
结构设计中必须控制在允许范围内
---
| 指标 | 说明 | 限值 | 依据 |
| 挠度 | 竖向位移 | L/250 | 规范 |
| $q$ | 均布荷载 | — | 材料力学 |
| $E$ | 弹性模量 | — | 材料力学 |
| $I$ | 惯性矩 | — | 材料力学 |
| $\\delta$ | 挠度值 | — | **四次方正比** |
| 对比项 | 旧方案 | 新方案 | |
| 刚度 | 低 | 高 | |
| 造价 | 高 | 低 | |
| 施工 | 难 | 易 | |
| 维护 | 复杂 | 简单 | |
---
> 挠度控制是结构刚度设计的核心结论
:::
"""


def test_no_page_number_placeholder_titles():
    artifact = extract_artifact(BAD_MARP, default_type="pptx")
    assert artifact.slides, "应解析出幻灯片"
    for s in artifact.slides:
        assert not s.title.startswith("第 "), f"禁止页码占位标题: {s.title}"
        assert not (s.title.startswith("第 ") and s.title.endswith("页"))


def test_cover_title_falls_back_to_artifact_title():
    artifact = extract_artifact(BAD_MARP, default_type="pptx")
    cover = artifact.slides[0]
    assert cover.title == "结构刚度关键指标——挠度（Deflection）深度解析"


def test_content_page_title_not_placeholder():
    artifact = extract_artifact(BAD_MARP, default_type="pptx")
    # 第 2 页无 # 标题：应落到首条短要点或「内容页 N」，绝不能是「第 2 页」
    s2 = artifact.slides[1]
    assert s2.title != "第 2 页"
    assert s2.title
    # 第 4 页 Quote：应用金句或「内容页 4」
    quote_page = artifact.slides[-1]
    assert quote_page.title != f"第 {quote_page.index} 页"


def test_quote_export_uses_real_quote_without_bullets():
    """运算符优先级修复：无 bullet 时也必须展示 quote_text。"""
    model = SlideModel(
        index=1,
        title="金句页",
        layout=SlideLayoutType.QUOTE,
        quote_text="挠度控制是结构刚度设计的核心结论",
    )
    artifact = ArtifactModel(
        artifact_type=ArtifactType.PPTX,
        title="金句测试",
        raw_content="",
        slides=[model],
    )
    from doc2mind.core.creator.themes import get_theme

    # 直接调用渲染逻辑：用 python-pptx 打开后检查文本
    from pptx import Presentation as PptPresentation
    from pptx.util import Inches

    prs = PptPresentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    theme = get_theme("tech_blue")
    exporter = PptxExporter()
    exporter._render_quote(slide, model, theme, 1)

    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            texts.append(shape.text_frame.text)
    joined = "\n".join(texts)
    assert "挠度控制是结构刚度设计的核心结论" in joined
    assert "核心结论与洞察" not in joined


def test_quote_bullet_fallback_still_works():
    model = SlideModel(
        index=1,
        title="金句页",
        layout=SlideLayoutType.QUOTE,
        quote_text="",
        bullet_points=["真实要点金句"],
    )
    from doc2mind.core.creator.themes import get_theme
    from pptx import Presentation as PptPresentation
    from pptx.util import Inches

    prs = PptPresentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    exporter = PptxExporter()
    exporter._render_quote(slide, model, get_theme("tech_blue"), 1)
    texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
    assert any("真实要点金句" in t for t in texts)


def test_validate_gate_blocks_empty_body_page():
    artifact = ArtifactModel(
        artifact_type=ArtifactType.PPTX,
        title="空壳",
        raw_content="",
        slides=[
            SlideModel(index=1, title="封面", subtitle="副标题", is_cover=True, layout=SlideLayoutType.COVER),
            SlideModel(index=2, title="空页", layout=SlideLayoutType.GENERAL),
        ],
    )
    errors = validate_pptx_export(artifact)
    assert errors
    assert any("空正文页" in e for e in errors)


def test_validate_gate_blocks_markdown_residue():
    artifact = ArtifactModel(
        artifact_type=ArtifactType.PPTX,
        title="残留",
        raw_content="",
        slides=[
            SlideModel(
                index=1,
                title="表格页",
                layout=SlideLayoutType.TABLE,
                table_data=[
                    ["$q$", "**四次方正比**", "$E I$", "$\\delta$"],
                    ["a", "b", "c", "d"],
                ],
            ),
        ],
    )
    errors = validate_pptx_export(artifact)
    assert errors
    assert any("Markdown" in e for e in errors)


def test_export_artifact_refuses_bad_structure(tmp_path: Path):
    """门禁：结构不合格时不写盘，返回 ok=False。"""
    out = tmp_path / "should_not_exist.pptx"
    res = export_artifact(
        content="""
:::artifact type="pptx" title="空壳导出"
---
---
---
""",
        target_format="pptx",
        output_path=out,
    )
    assert res.ok is False
    assert res.error
    assert "结构不合格" in res.error
    assert not out.exists()


def test_export_artifact_accepts_good_deck(tmp_path: Path):
    out = tmp_path / "good.pptx"
    res = export_artifact(
        content="""
:::artifact type="pptx" title="合格汇报" theme="tech_blue"
---
# 合格汇报
## 副标题说明
---
# 核心结论
- 挠度限值取 L/250
- 需同时满足强度与刚度
---
# 对比数据
| 维度 | 方案A | 方案B |
| 造价 | 高 | 低 |
| 工期 | 长 | 短 |
---
# 总结
> 刚度与造价需协同权衡
:::
""",
        target_format="pptx",
        output_path=out,
    )
    assert res.ok is True, res.error
    assert out.exists()
    prs = Presentation(str(out))
    assert len(prs.slides) == 4
    titles = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t and len(t) < 40:
                    titles.append(t)
                    break
    # 不允许出现页码占位标题
    assert not any(t.startswith("第 ") and t.endswith("页") for t in titles)
